"""
mcp_document_processing.py
FastMCP server exposing a Document‑processing tool
that works without the camel package.
"""

# --------------------------------------------------------------------------- #
#  Imports
# --------------------------------------------------------------------------- #
import asyncio, os, io, json, subprocess
from typing import Tuple, Optional, List, Literal
from pathlib import Path

from loguru import logger
from retry import retry

from mcp.server.fastmcp import FastMCP
import anyio

# --- your own helper toolkits ------------------------------------------------ #
#   (provide these scripts in the Python path)
from image_tool import ask_question_about_image
from excel_tool import ExcelToolkit
from video_tool import ask_question_about_video
# --- third‑party libs already used ------------------------------------------ #
import assemblyai as aai
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx import Presentation
from PIL import Image
from docx2markdown._docx_to_markdown import docx_to_markdown
from chunkr_ai import Chunkr
import xmltodict
import nest_asyncio
nest_asyncio.apply()

from dotenv import load_dotenv
# --- 修改开始 ---
dotenv_path = Path(__file__).resolve().parent.parent / "client" / ".env"
load_dotenv(dotenv_path=dotenv_path)
# --- 修改结束 ---

# --------------------------------------------------------------------------- #
#  Toolkit implementation (no camel.BaseToolkit!)
# --------------------------------------------------------------------------- #
class DocumentProcessingToolkit:
    """
    This tool exposes a **general‑purpose document‑processing endpoint** that
    converts almost any common file you point it to into **clean, readable
    text or Markdown**.  It is useful whenever an agent needs to “look inside”
    an arbitrary file before reasoning over its contents.
    """

    def __init__(self, cache_dir: Optional[str] = None):
        self.excel_tool = ExcelToolkit()
        self.cache_dir = cache_dir or "tmp/"

    # --------------------------------------------------------------------- #
    #  Public façade
    # --------------------------------------------------------------------- #
    @retry(Exception,tries=5, delay=2, backoff=2)
    def extract_document_content(self, document_path: str) -> Tuple[bool, str]:
        logger.debug(f"[extract_document_content] {document_path=}")

        # 1. Images ----------------------------------------------------------------
        if document_path.lower().endswith((".jpg", ".jpeg", ".png")):
            caption = asyncio.run(
                ask_question_about_image(
                    document_path,
                    "Please make a detailed caption about the image."
                )
            )
            return True, caption

        # 2. Audio -----------------------------------------------------------------
        if document_path.lower().endswith((".mp3", ".wav", ".m4a")):
            aai.settings.api_key = os.getenv("ASSEMBLYAI_API_KEY")
            if not aai.settings.api_key:
                return False, "ASSEMBLYAI_API_KEY not set."
            config = aai.TranscriptionConfig(speech_model=aai.SpeechModel.best)
            transcript = aai.Transcriber(config=config).transcribe(document_path)
            logger.info(transcript.text)
            if transcript.status == "error":
                raise RuntimeError(f"Transcription failed: {transcript.error}")
            return True, transcript.text or ""

        # 3. PPTX ------------------------------------------------------------------
        if document_path.lower().endswith(".pptx"):
            return True, asyncio.run(self._extract_pptx(document_path))

        # 4. Spreadsheets -----------------------------------------------------------
        if document_path.lower().endswith((".xls", ".xlsx", ".csv")):
            return True, self.excel_tool.extract_excel_content(document_path)

        # 5. Zip --------------------------------------------------------------------
        if document_path.lower().endswith(".zip"):
            return True, f"The extracted files are: {self._unzip_file(document_path)}"

        # 6. Simple text‑like formats ----------------------------------------------
        simple_readers = {
            ".py":  lambda p: open(p, encoding="utf‑8").read(),
            ".txt": lambda p: open(p, encoding="utf‑8").read(),
        }
        if any(document_path.lower().endswith(ext) for ext in simple_readers):
            reader = simple_readers[os.path.splitext(document_path)[1]]
            return True, reader(document_path)

        # 7. JSON                                                                   #
        if document_path.lower().endswith((".json", ".jsonl", ".jsonld")):
            return True, self._extract_json(document_path, encoding="utf‑8")
        

        # 8. XML                                                                    #
        if document_path.lower().endswith(".xml"):
            data = open(document_path, encoding="utf‑8").read()
            try:
                # Convert XML to dict and then pretty-print as JSON string
                return True, json.dumps(xmltodict.parse(data), indent=4)
            except Exception:
                return True, data

        # 9. DOCX → markdown -------------------------------------------------------
        if document_path.lower().endswith(".docx"):
            md_path = f"{os.path.basename(document_path)}.md"
            docx_to_markdown(document_path, md_path)
            return True, open(md_path, encoding="utf‑8").read()

        # 10. MOV video ------------------------------------------------------------
        if document_path.lower().endswith((".mov", ".mp4", ".avi")):
            description = asyncio.run(ask_question_about_video(
                document_path, "Please make a detailed description about the video."
            ))
            return True, description

        # 11. Fallback – Chunkr / PDF text -----------------------------------------
        return self._try_chunkr_then_fallback(document_path)

    # ------------------------------------------------------------------------- #
    #  helpers
    # ------------------------------------------------------------------------- #
    def _extract_json(self, json_path: str, encoding: str = "utf‑8") -> str:
        with open(json_path, 'r', encoding=encoding) as f:
            if json_path.lower().endswith((".json",".jsonld")):
                return json.dumps(json.load(f), indent=4)
            elif json_path.lower().endswith(".jsonl"):
                return json.dumps([json.loads(line) for line in f], indent=4)
        return ""                    

    async def _extract_pptx(self, pptx_path: str) -> str:
        prs = Presentation(pptx_path)
        base = pptx_path.rsplit(".", 1)[0]
        out = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            txt = [f"Page {slide_idx}"]
            captions = []
            img_count = 0

            for shape_idx, shape in enumerate(slide.shapes):
                shape_text = getattr(shape, "text", None)
                if shape_text and shape_text.strip():
                    txt.append(shape_text.strip())

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    img_count += 1
                    # --- 修改开始 ---
                    shape_image = getattr(shape, "image", None)
                    if shape_image:
                        img = Image.open(io.BytesIO(shape_image.blob))
                        img_path = f"{base}_slide_{slide_idx}_img_{shape_idx}.png"
                        img.save(img_path)
                        captions.append(
                            f"Image {img_count}: "
                            + await ask_question_about_image(
                                img_path, "Please make a detailed caption about the image."
                            )
                        )
                    # --- 修改结束 ---

            out.append("\n".join(txt + captions))

        return "\n\n".join(out)

    def _try_chunkr_then_fallback(self, path: str) -> Tuple[bool, str]:
        # Try Chunkr first if API key is available
        if os.getenv("CHUNKR_API_KEY"):
            try:
                text = asyncio.run(
                    self._extract_with_chunkr(path, output_format="markdown")
                ) # type: ignore
                return True, text
            except Exception as e:
                logger.warning(f"Chunkr failed: {e}")

        # Fallback for PDF
        if path.lower().endswith(".pdf"):
            try:
                from PyPDF2 import PdfReader
                text = "".join(
                    p.extract_text() for p in PdfReader(open(path, "rb")).pages
                )
                return True, text
            except Exception as e2:
                return False, f"PDF fallback failed: {e2}"
        
        return False, f"Unsupported file type or processing error for: {path}"

    async def _extract_with_chunkr(
        self, path: str, output_format: Literal["json", "markdown"] = "markdown"
    ) -> str:
        chunkr = Chunkr(api_key=os.getenv("CHUNKR_API_KEY"))
        result = await chunkr.upload(path) # type: ignore

        if result.status == "Failed":
            raise RuntimeError(result.message)

        out_path = f"{os.path.basename(path)}.{ 'json' if output_format=='json' else 'md' }"
        (result.json if output_format == "json" else result.markdown)(out_path)
        return open(out_path, encoding="utf‑8").read()

    def _unzip_file(self, zip_path: str) -> List[str]:
        dest = os.path.join(self.cache_dir, os.path.splitext(os.path.basename(zip_path))[0])
        os.makedirs(dest, exist_ok=True)
        subprocess.run(["unzip", "-o", zip_path, "-d", dest], check=True)
        return [os.path.join(r, f) for r, _, fs in os.walk(dest) for f in fs]


# --------------------------------------------------------------------------- #
#  FastMCP server
# --------------------------------------------------------------------------- #
mcp = FastMCP("document_processing")
toolkit = DocumentProcessingToolkit()


@mcp.tool()
async def process_document(document_path: str) -> str:
    """
    Process a document at the given *document_path*. The document can be multimedia
    (image, audio, video), a presentation (PPTX), a spreadsheet, a ZIP archive,
    a text file, JSON, XML, Word document, or PDF.
   Return the extracted text / markdown representation of *document_path*.
    """
    success, content = await anyio.to_thread.run_sync( # type: ignore
        toolkit.extract_document_content, document_path
    ) # type: ignore
    if not success:
        raise ValueError(content)
    return content


# --------------------------------------------------------------------------- #
#  Entrypoint
# --------------------------------------------------------------------------- #
if __name__ == "__main__":
    mcp.run(transport="stdio")

